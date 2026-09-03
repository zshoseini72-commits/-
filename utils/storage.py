import logging
import os
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from config import file_name, file_name_member

logger = logging.getLogger(__name__)

class Storage:

    @classmethod
    def check_file(cls, file: str) -> bool:

        logger.debug(f"بررسی وجود فایل: {file}")

        if not os.path.exists(file):
            logger.warning(f"فایل پیدا نشد: {file}")
            print("هنوز هیچ داده ای ثبت نشده است")
            return False

        logger.info(f"فایل پیدا شد: {file}")
        return True


    @staticmethod
    def workbook() -> tuple[Workbook, Worksheet]:

        logger.debug("در حال باز کردن فایل کتاب‌ها")

        try:
            wb = load_workbook(file_name)
            ws = wb.active

            if ws is None:
                logger.error("Worksheet فایل کتاب‌ها پیدا نشد")
                raise ValueError("Worksheet پیدا نشد")

            logger.info("فایل کتاب‌ها با موفقیت باز شد")

            return wb, ws

        except PermissionError as error:
            logger.error(f"دسترسی به فایل کتاب‌ها امکان‌پذیر نیست: {error}")
            raise

        except FileNotFoundError as error:
            logger.error(f"فایل کتاب‌ها پیدا نشد: {error}")
            raise

        except Exception as error:
            logger.critical(f"خطای غیرمنتظره در باز کردن فایل کتاب‌ها: {error}")
            raise


    @staticmethod
    def load_sheet() -> tuple[Workbook, Worksheet]:

        logger.debug("در حال باز کردن فایل اعضا")

        try:
            wb = load_workbook(file_name_member)
            ws = wb.active

            if ws is None:
                logger.error("Worksheet فایل اعضا پیدا نشد")
                raise ValueError("Worksheet پیدا نشد")

            logger.info("فایل اعضا با موفقیت باز شد")

            return wb, ws

        except PermissionError as error:
            logger.error(f"دسترسی به فایل اعضا امکان‌پذیر نیست: {error}")
            raise

        except FileNotFoundError as error:
            logger.error(f"فایل اعضا پیدا نشد: {error}")
            raise

        except Exception as error:
            logger.critical(f"خطای غیرمنتظره در باز کردن فایل اعضا: {error}")
            raise


class ErrorCodeMeli(Exception):

    def __init__(self, code):
        self.code = code
        super().__init__()

    def __str__(self):
        return "کد ملی را صحیح وارد کنید"