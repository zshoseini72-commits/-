import logging
import os
from plistlib import InvalidFileException
from typing import Generator
from openpyxl.cell import Cell, MergedCell
from openpyxl import load_workbook, Workbook
from config import file_name
from models.book import Book
from utils.decorators import timer, title
from services.service_static import ServiceStatic
from utils.storage import Storage
from utils.contracts import BaseService, RepositoryProtocol, SimpleRepository

logger = logging.getLogger(__name__)


class BookService(BaseService):

    # تمرین Protocol + Generic: هر Repository که قرارداد add(Book) را رعایت کند قابل استفاده است.
    repository: RepositoryProtocol[Book] = SimpleRepository[Book]()


    @classmethod
    def service_name(cls) -> str:
        return "BookService"

    @classmethod
    def book_generator(cls) -> Generator[tuple[Cell | MergedCell, ...], None, None]:

        logger.debug("شروع خواندن کتاب‌ها")

        if not Storage.check_file(file_name):
            logger.warning("فایل کتاب‌ها وجود ندارد")
            print("هنوز هیچ کتابی ثبت نشده است")
            return

        wb = None

        try:
            wb, ws = Storage.workbook()

            logger.info("فایل کتاب‌ها برای خواندن باز شد")

            for row in ws.iter_rows(min_row=2):
                logger.debug(f"خواندن ردیف کتاب: {row[3].value}")
                yield row

        except Exception as error:
            logger.error(f"خطا هنگام خواندن کتاب‌ها: {error}")
            raise

        finally:
            if wb is not None:
                wb.close()
                logger.debug("فایل کتاب‌ها بسته شد")


    @classmethod
    def find_book(cls, put_id: int) -> str:

        logger.info(f"جستجوی کتاب با ID: {put_id}")

        for row in cls.book_generator():

            try:
                if row[3].value == put_id:
                    logger.info(f"کتاب با ID {put_id} پیدا شد")

                    print("کتاب پیدا شد")
                    return ServiceStatic.print_book(row)

            except ValueError as error:
                logger.warning(
                    f"ID وارد شده برای جستجوی کتاب نامعتبر است: {put_id}"
                )
                logger.error(f"جزئیات خطا: {error}")

                return "آی دی صحیح نمیباشد"

        logger.warning(f"کتابی با ID {put_id} پیدا نشد")

        return "کتاب موجود نمی باشد"


    @classmethod
    @timer
    @title
    def show_books(cls) -> None:

        logger.info("نمایش لیست کتاب‌ها شروع شد")

        for row in cls.book_generator():

            try:
                print(ServiceStatic.print_book(row))

            except Exception as error:
                logger.error(
                    f"خطا هنگام نمایش اطلاعات کتاب: {error}"
                )
                print(error, "خطا: ")

        logger.info("نمایش لیست کتاب‌ها تمام شد")


    @classmethod
    def add_book(cls, book: Book) -> str:

        logger.info(f"شروع ثبت کتاب: {book.title}")

        workbook = None

        try:

            if os.path.exists(file_name):

                logger.debug("فایل کتاب‌ها وجود دارد؛ در حال باز کردن فایل")

                workbook = load_workbook(file_name)
                sheet = workbook.worksheets[0]

            else:

                logger.warning(
                    "فایل کتاب‌ها وجود ندارد؛ فایل جدید ساخته می‌شود"
                )

                workbook = Workbook()
                sheet = workbook.worksheets[0]

                sheet.append(
                    ["عنوان", "نویسنده", "سال", "ID", "وضعیت"]
                )

            sheet.append(
                [
                    book.title,
                    book.author,
                    book.year,
                    book.id,
                    book.state.value
                ]
            )

            workbook.save(file_name)

            # ثبت در Repository عمومی (Generic) برای تمرین معماری؛
            # ذخیره اصلی و دائمی همچنان Excel است.
            cls.repository.add(book)

            logger.info(
                f"کتاب با موفقیت ثبت شد - ID: {book.id}"
            )

            print(
                f" عنوان: {book.title} "
                f"نویسنده: {book.author} "
                f"سال انتشار: {book.year} "
                f"آی دی: {book.id} "
                f"وضعیت: {book.state.value}"
            )

            return "کتاب ثبت شد"

        except Exception as error:

            logger.error(
                f"ثبت کتاب با خطا مواجه شد - "
                f"ID: {book.id} - خطا: {error}"
            )

            print(error)

            return " ثبت کتاب با خطا مواجه شد "

        finally:

            if workbook is not None:
                workbook.close()
                logger.debug("فایل کتاب‌ها بسته شد")


    @classmethod
    def remove_book(cls, put_id: int) -> str:

        logger.info(f"شروع حذف کتاب با ID: {put_id}")

        if not Storage.check_file(file_name):

            logger.warning("فایل کتاب‌ها برای حذف وجود ندارد")

            return " فایلی موجود نیست"

        wb = None

        try:

            wb, ws = Storage.workbook()

            for index, row in enumerate(
                ws.iter_rows(min_row=2),
                start=2
            ):

                if row[3].value == put_id:

                    logger.info(
                        f"کتاب با ID {put_id} پیدا شد و آماده حذف است"
                    )

                    print("کتاب پیدا شد")
                    print(ServiceStatic.print_book(row))

                    ws.delete_rows(index)
                    wb.save(file_name)

                    logger.info(
                        f"کتاب با ID {put_id} با موفقیت حذف شد"
                    )

                    return "کتاب حذف شد"

            logger.warning(
                f"کتابی با ID {put_id} برای حذف پیدا نشد"
            )

            return "آی دی موجود نمی باشد"

        except Exception as error:

            logger.error(
                f"حذف کتاب با ID {put_id} با خطا مواجه شد: {error}"
            )

            print("خطا: ", error)

            return "حذف کتاب با خطا مواجه شد"

        finally:

            if wb is not None:
                wb.close()
                logger.debug("فایل کتاب‌ها بسته شد")


    @classmethod
    def edit_book(
        cls,
        put_id: int,
        choice: str,
        change: str
    ) -> None:

        logger.info(
            f"شروع ویرایش کتاب - ID: {put_id}"
        )

        if not Storage.check_file(file_name):

            logger.warning(
                "فایل کتاب‌ها برای ویرایش وجود ندارد"
            )

            return

        wb = None

        try:

            wb, ws = Storage.workbook()

            for index, row in enumerate(
                ws.iter_rows(min_row=2),
                start=2
            ):

                if row[3].value == put_id:

                    logger.info(
                        f"کتاب با ID {put_id} پیدا شد"
                    )

                    print(ServiceStatic.print_book(row))
                    print("کتاب پیدا شد")

                    if choice == "1":

                        logger.debug(
                            f"تغییر عنوان کتاب با ID {put_id}"
                        )

                        row[0].value = change

                    elif choice == "2":

                        logger.debug(
                            f"تغییر نویسنده کتاب با ID {put_id}"
                        )

                        row[1].value = change

                    elif choice == "3":

                        logger.debug(
                            f"تغییر سال کتاب با ID {put_id}"
                        )

                        row[2].value = int(change)

                    else:

                        logger.warning(
                            f"انتخاب نامعتبر برای ویرایش کتاب: {choice}"
                        )

                        print("انتخاب نامعتبر است.")
                        return

                    wb.save(file_name)

                    logger.info(
                        f"ویرایش کتاب با ID {put_id} با موفقیت انجام شد"
                    )

                    print(ServiceStatic.print_book(row))
                    print("تغییرات با موفقیت اعمال شد")

                    return

            logger.warning(
                f"کتابی با ID {put_id} برای ویرایش پیدا نشد"
            )

            print("آی دی صحیح نمیباشد")

        except (
            InvalidFileException,
            PermissionError,
            FileNotFoundError
        ) as error:

            logger.error(
                f"خطا در فایل هنگام ویرایش کتاب "
                f"با ID {put_id}: {error}"
            )

            print("خطا در فایل:", error)
            raise

        except ValueError as error:

            logger.warning(
                f"سال وارد شده برای کتاب با ID {put_id} نامعتبر است: {error}"
            )

            raise

        finally:

            if wb is not None:
                wb.close()
                logger.debug("فایل کتاب‌ها بسته شد")


