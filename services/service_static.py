import logging
import os
from openpyxl import load_workbook
from openpyxl.cell import MergedCell, Cell

from config import file_name, file_name_member


logger = logging.getLogger(__name__)


class ServiceStatic:

    @staticmethod
    def count_rows(file) -> int:
        logger.debug("شروع شمارش ردیف‌ها")

        if not os.path.exists(file):
            logger.warning("فایل وجود ندارد: %s", file)
            return 0

        wb=None
        try:
            wb = load_workbook(file)

            logger.debug("فایل باز شد: %s", file)

            return wb.active.max_row - 1

        finally:
            if wb is not None:
                wb.close()
                logger.debug("فایل بسته شد: %s", file)

    @staticmethod
    def count_books() -> int:
        logger.debug("شروع شمارش کتاب‌ها")
        return ServiceStatic.count_rows(file_name)

    @staticmethod
    def count_members() -> int:
        logger.debug("شروع شمارش اعضا")
        return ServiceStatic.count_rows(file_name_member)

    @property
    def books(self):
        logger.debug("شروع دریافت تعداد کتاب‌ها")
        return self.count_books()

    @property
    def members(self):
        logger.debug("شروع دریافت تعداد اعضا")
        return self.count_members()

    def __len__(self) -> int:
        logger.debug("شروع محاسبه مجموع کتاب‌ها و اعضا")
        return self.books + self.members

    def __str__(self) -> str:
        logger.debug("شروع نمایش تعداد کتاب‌ها، اعضا و مجموع")

        return (
            f"تعداد کل کتاب‌ها: {self.books}\n"
            f"تعداد کل اعضا: {self.members}\n"
            f"مجموع: {len(self)}"
        )

    @staticmethod
    def print_book(row: tuple[Cell | MergedCell, ...]) -> str:
        logger.debug(
            "نمایش اطلاعات کتاب با شناسه: %s",
            row[3].value
        )

        return (
            f"عنوان: {row[0].value}"
            f" - نویسنده: {row[1].value}"
            f" - سال انتشار: {row[2].value}"
            f" - شناسه: {row[3].value}"
            f" - وضعیت: {row[4].value}"
        )

    @staticmethod
    def print_member(row: tuple[Cell | MergedCell, ...]) -> str:
        logger.debug(
            "نمایش اطلاعات عضو با کد ملی: %s",
            row[3].value
        )

        return (
            f"نام: {row[0].value}"
            f" نام خانوادگی: {row[1].value}"
            f" تلفن: {row[2].value}"
            f" کد ملی: {row[3].value}"
            f" عضویت: {row[4].value}"
            f" شهر: {row[5].value}"
            f" آدرس: {row[6].value}"
            f" ایمیل: {row[7].value}"
        )