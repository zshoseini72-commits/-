import logging
from plistlib import InvalidFileException
from typing import Generator
from openpyxl.cell import Cell, MergedCell
import os
from models.member import Member
from openpyxl import Workbook, load_workbook
from services.service_static import file_name_member, ServiceStatic
from utils import BaseService ,Storage ,timer


logger = logging.getLogger(__name__)

class MemberService(BaseService):

    @classmethod
    def service_name(cls) -> str:
        return "خدمات بخش اعضا"

    @classmethod
    def member_generator(cls) -> Generator[tuple[Cell | MergedCell, ...], None, None]:

        logger.debug("شروع خواندن فایل اعضا")

        if not Storage.check_file(file_name_member):
            logger.warning("فایل اعضا وجود ندارد")
            return

        wb=None
        try:
            wb, ws = Storage.load_sheet()

            logger.info("فایل اعضا برای خواندن باز شد")

            for row in ws.iter_rows(min_row=2):
                logger.info(f"خواندن ردیف عضو با کد ملی {row[3].value}")
                yield row

        except (PermissionError, FileNotFoundError):
            logger.exception(f" خطا هنگام خواندن اعضا")
            raise

        finally:
            if wb is not None:
                wb.close()
                logger.info("فایل اعضا بسته شد")

    @classmethod
    def find_member(cls, code_meli: str) -> str:
        logger.debug(f"جستجوی عضو با کد ملی {code_meli}")

        try:
            for row in cls.member_generator():

                if row[3].value == code_meli:
                    logger.info(f"کاربر با کد ملی {code_meli} پیدا شد")

                    print(ServiceStatic.print_member(row))

                    return "کاربر پیدا شد"

        except Exception:
            logger.exception(
                f"هنگام جستجوی عضو با کد ملی {code_meli} خطا رخ داد: "
            )
            raise

        logger.warning(f"عضو با کد ملی {code_meli} موجود نبود")
        return "کاربر موجود نمی‌باشد"



    @classmethod
    @timer
    def show_members(cls) -> None:

        logger.debug("نمایش لیست اعضا")

        try:
            for row in cls.member_generator():
                print(ServiceStatic.print_member(row))

        except Exception:

            logger.exception(f" هنگام نمایش اعضا خطا رخ داد ")
            raise

        logger.info("نمایش لیست لعضا پایان یافت")


    @classmethod
    def add_member(cls, member: Member) -> str:

        logger.debug("شروع فرایند اضافه کردن عضو")
        workbook=None

        try:
            if os.path.exists(file_name_member):
                logger.info("فایل اعضا وجود دارد و در حال باز کردن فایل اعضا")

                workbook = load_workbook(file_name_member)
                sheet = workbook.worksheets[0]
            else:
                logger.info("فایل اعضا وجود ندارد- فایل جدید ساخته میشود")

                workbook = Workbook()
                sheet = workbook.worksheets[0]
                # ساخت عنوان ستون‌ها

                logger.debug("هدر - ردیف اول فایل اعضا ساخته میشود")
                sheet.append(["نام", "نام خانوادگی", "تلفن", "کد ملی", "عضویت","شهر", "آدرس", "ایمیل"])

            logger.debug("عضو جدید اضافه میشود")
            # ذخیره اطلاعات در یک ردیف اکسل
            sheet.append([member.firstname, member.lastname, member.phone, member.code_meli, member.role.value, member.city,
                          member.address, member.email])
            workbook.save(file_name_member)
            logger.info(f"عضو جدید با کد ملی {member.code_meli} اضافه شد")

            return (
                f" نام: {member.firstname}"
                f" نام خانوادگی: {member.lastname}"
                f" تلفن: {member.phone}"
                f" کد ملی: {member.code_meli}"
                f" عضویت: {member.role.value}"
                f" شهر: {member.city}"
                f" آدرس : {member.address}"
                f" ایمیل : {member.email}"
                "کابر ثبت شد"
            )


        except PermissionError as e:

            logger.exception(f"در ثبت عضو با کد ملی {member.code_meli} خطا رخ داد")
            return f"دسترسی به فایل امکان‌پذیر نیست: {e}"

        except InvalidFileException as e:
            logger.exception(f"در ثبت عضو با کد ملی {member.code_meli} خطا رخ داد")
            return f" خطا فایل Excel نامعتبر است: {e}"

        finally:
            if workbook is not None:
                workbook.close()
                logger.debug("فایل اعضا بسته شد")



    @classmethod
    def remove_member(cls, code_meli: str) -> None:
        logger.debug(f"شروع فرایند خواندن فایل اعضا جهت حذف عضو با کد ملی {code_meli}")

        if not Storage.check_file(file_name_member):
            logger.warning("فایل اعضا وجود ندارد")
            return

        wb= None

        try:
            wb, ws = Storage.load_sheet()
            logger.debug("فایل اعضا باز شد")


            found = False
            # جستجوی کاربر در ردیف‌ها
            for index, row in enumerate(ws.iter_rows(min_row=2), start=2):

                if row[3].value == code_meli:

                    logger.info(f"کاربر با کد ملی {code_meli} پیدا شد و آماده حذف است ")
                    found = True
                    print("کاربر پیدا شد")
                    print(ServiceStatic.print_member(row))

                    ws.delete_rows(index)
                    wb.save(file_name_member)
                    print("کاربر فوق حذف شد")
                    logger.info(f"کاربر با کد ملی {code_meli} حذف شد")

                    break

            if not found:
                logger.warning(f"کاربر با کد ملی {code_meli} موجود نیست ")
                print("کاربر موجود نمی باشد")

        except PermissionError as e:
            logger.exception(f"در حذف عضو با کد ملی {code_meli} خطا رخ داد")
            print( f" خطا: {e}")

        except InvalidFileException as e:
            logger.exception(f"در حذف عضو با کد ملی {code_meli} خطا رخ داد")
            print(f" خطا: {e}")

        finally:
            if wb is not None:
                wb.close()
                logger.debug("فایل اعضا بسته شد")

    @classmethod
    def edit_member(cls, code_meli: str, choice: str, change: str) -> None:
        logger.debug(f"شروع فرایند ویرایش عضو با کد ملی {code_meli}")

        if not Storage.check_file(file_name_member):
            logger.warning("فایل اعضا وجود ندارد")
            return

        wb = None

        try:
            logger.debug("فایل اعضا باز می‌شود")
            wb, ws = Storage.load_sheet()

            found = False

            for row in ws.iter_rows(min_row=2):

                if row[3].value == code_meli:
                    logger.info(f"عضو با کد ملی {code_meli} پیدا شد")

                    found = True

                    print(ServiceStatic.print_member(row))
                    print("کاربر پیدا شد")

                    try:
                        column_index = int(choice) - 1
                    except ValueError:
                        logger.warning(f"مقدار ورودی برای ویرایش عضو نامعتبر است: {choice}")
                        print("ورودی نامعتبر است.")
                        return

                    if not 0 <= column_index <= 7 or column_index == 3:
                        logger.warning(f"انتخاب نامعتبر برای ویرایش عضو: {choice}")
                        print("انتخاب نامعتبر است.")
                        return

                    row[column_index].value = change

                    logger.debug(
                        f"عضو با کد ملی {code_meli} تغییر کرد؛ "
                        f" مقدار جدید: {change}"
                    )

                    print(ServiceStatic.print_member(row))
                    print("تغییرات با موفقیت اعمال شد")

                    wb.save(file_name_member)

                    logger.info(
                        f"تغییرات عضو با کد ملی {code_meli} با موفقیت ذخیره شد"
                    )

                    break

            if not found:
                logger.warning(f"کاربر با کد ملی {code_meli} موجود نیست")
                print("کاربر موجود نمی‌باشد")

        except PermissionError as e:
            logger.exception(
                f"در ویرایش عضو با کد ملی {code_meli} خطای دسترسی رخ داد"
            )
            print("خطا:", e)

        except InvalidFileException as e:
            logger.exception(
                f"در ویرایش عضو با کد ملی {code_meli} فایل Excel نامعتبر است"
            )
            print("خطا:", e)

        finally:
            if wb is not None:
                wb.close()
                logger.debug("فایل اعضا بسته شد")
